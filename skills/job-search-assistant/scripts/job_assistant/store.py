"""持久化行动账本：先登记，再操作，凭证确认后才标记成功。"""
import hashlib
import json
import shutil
import sqlite3
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from .inputs import web_url


def now():
    return datetime.now(timezone.utc).isoformat()


def packed(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


class Store:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(self.root / 'records.sqlite3', timeout=10)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute('PRAGMA foreign_keys=ON')
        self.conn.executescript('''
        CREATE TABLE IF NOT EXISTS runs (
          id TEXT PRIMARY KEY, config TEXT NOT NULL, approval TEXT NOT NULL,
          status TEXT NOT NULL, created TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS attempts (
          id TEXT PRIMARY KEY, run_id TEXT REFERENCES runs(id), company TEXT, url TEXT,
          job TEXT, channel TEXT, job_key TEXT, status TEXT, reason TEXT, started TEXT, updated TEXT);
        CREATE TABLE IF NOT EXISTS run_attempts (
          run_id TEXT REFERENCES runs(id), attempt_id TEXT REFERENCES attempts(id), attached TEXT,
          PRIMARY KEY(run_id,attempt_id));
        CREATE TABLE IF NOT EXISTS fields (
          attempt_id TEXT REFERENCES attempts(id), name TEXT, value TEXT, source TEXT,
          evidence_id TEXT, PRIMARY KEY(attempt_id,name));
        CREATE TABLE IF NOT EXISTS evidence (
          id TEXT PRIMARY KEY, attempt_id TEXT REFERENCES attempts(id), kind TEXT,
          path TEXT, sha256 TEXT, created TEXT);
        CREATE TABLE IF NOT EXISTS actions (
          id TEXT PRIMARY KEY, run_id TEXT REFERENCES runs(id), attempt_id TEXT REFERENCES attempts(id),
          kind TEXT, recipient TEXT, payload TEXT, dedup_key TEXT, status TEXT,
          evidence_id TEXT, reason TEXT, created TEXT, updated TEXT);
        CREATE INDEX IF NOT EXISTS action_dedup ON actions(dedup_key,status);
        CREATE TABLE IF NOT EXISTS conversations (
          id TEXT PRIMARY KEY, status TEXT, reason TEXT, latest_message TEXT, updated TEXT);
        CREATE TABLE IF NOT EXISTS messages (
          conversation_id TEXT, message_key TEXT, text TEXT, created TEXT,
          PRIMARY KEY(conversation_id,message_key));
        CREATE TABLE IF NOT EXISTS attention (
          id TEXT PRIMARY KEY, subject TEXT, reason TEXT, resolved INTEGER DEFAULT 0, created TEXT,
          notified INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS panel_commands (
          id TEXT PRIMARY KEY, run_id TEXT REFERENCES runs(id), action TEXT,
          status TEXT, error TEXT, created TEXT, updated TEXT);
        CREATE TABLE IF NOT EXISTS executor_activity (
          run_id TEXT PRIMARY KEY REFERENCES runs(id), operation TEXT, phase TEXT,
          updated TEXT, fresh_after REAL);
        CREATE TABLE IF NOT EXISTS attention_notes (
          id TEXT PRIMARY KEY, attention_id TEXT REFERENCES attention(id),
          message_key TEXT, text TEXT, created TEXT);
        ''')
        if 'notified' not in {row[1] for row in self.conn.execute('PRAGMA table_info(attention)')}:
            self.conn.execute('ALTER TABLE attention ADD COLUMN notified INTEGER DEFAULT 0')
        self.conn.commit()

    def close(self):
        self.conn.close()

    def one(self, table, identity):
        if table not in ('runs', 'attempts', 'actions', 'evidence', 'conversations'):
            raise ValueError('未知记录类型')
        row = self.conn.execute(f'SELECT * FROM {table} WHERE id=?', (identity,)).fetchone()
        if row is None:
            raise ValueError(f'{table} 记录不存在')
        return dict(row)

    def create_run(self, config, approval):
        if not str(approval).strip():
            raise ValueError('必须记录用户明确的本轮授权')
        if not config.get('channels') or not set(config['channels']) <= {'web', 'boss'}:
            raise ValueError('渠道配置无效')
        if 'boss' in config['channels'] and (type(config.get('max_contacts')) is not int or config['max_contacts'] < 1):
            raise ValueError('必须明确沟通上限')
        identity = uuid.uuid4().hex
        with self.conn:
            self.conn.execute('INSERT INTO runs VALUES (?,?,?,?,?)',
                              (identity, packed(config), approval, 'running', now()))
        return identity

    def check_run(self, run_id, kind=None):
        run = self.one('runs', run_id)
        if run['status'] not in ('running', 'watching'):
            raise ValueError('运行已暂停或停止')
        if run['status'] == 'watching' and kind in ('greet', 'submit', 'fill'):
            raise ValueError('仅值守模式已暂停新增投递')
        return json.loads(run['config'])

    def control(self, run_id, status):
        if status not in ('running', 'watching', 'paused', 'stopped'):
            raise ValueError('无效运行状态')
        with self.conn:
            cursor = self.conn.execute("UPDATE runs SET status=? WHERE id=? AND status!='stopped'", (status, run_id))
            if cursor.rowcount != 1:
                raise ValueError('运行不存在或已停止；请新建授权轮次')
            if status in ('running', 'watching'):
                self.conn.execute('''INSERT INTO executor_activity VALUES (?, '', '', NULL, ?)
                  ON CONFLICT(run_id) DO UPDATE SET fresh_after=excluded.fresh_after''', (run_id, time.time()))
            if status in ('paused', 'stopped'):
                self.conn.execute("UPDATE panel_commands SET status='cancelled',updated=? WHERE run_id=? AND status='pending'", (now(), run_id))

    def begin_attempt(self, run_id, company, url, job, channel, job_key):
        with self.conn:
            self.conn.execute('BEGIN IMMEDIATE')
            return self._begin_attempt(run_id, company, url, job, channel, job_key)

    def _begin_attempt(self, run_id, company, url, job, channel, job_key):
        config = self.check_run(run_id, 'fill' if channel == 'web' else 'greet')
        if channel not in config['channels']:
            raise ValueError('渠道不在本轮授权范围内')
        company_scope = config.get('companies' if channel == 'web' else 'boss_companies')
        if company_scope and company not in company_scope:
            raise ValueError('公司不在本轮授权范围内')
        web_url(url)
        if not all(str(v).strip() for v in (company, job, job_key)):
            raise ValueError('公司、岗位和稳定岗位标识不能为空')
        existing = self.conn.execute('''SELECT id FROM attempts WHERE channel=? AND
            (job_key=? OR (?='web' AND company=?)) AND status IN ('in_progress','success','uncertain','needs_user')''',
            (channel, job_key, channel, company)).fetchone()
        if existing:
            raise ValueError(f'重复投递或结果待核实：{existing[0]}')
        identity, stamp = uuid.uuid4().hex, now()
        with self.conn:
            self.conn.execute('INSERT INTO attempts VALUES (?,?,?,?,?,?,?,?,?,?,?)',
                              (identity, run_id, company, url, job, channel, job_key, 'in_progress', '', stamp, stamp))
        return identity

    def set_attempt(self, attempt_id, status, reason=''):
        item = self.one('attempts', attempt_id)
        if item['status'] == 'success':
            raise ValueError('已成功的申请不能降级以重新投递')
        unresolved = self.conn.execute("SELECT 1 FROM actions WHERE attempt_id=? AND status IN ('pending','executing','uncertain')", (attempt_id,)).fetchone()
        if unresolved and status not in ('uncertain', 'needs_user'):
            raise ValueError('存在结果待核实行动，先核验行动结果')
        if self.conn.execute("SELECT 1 FROM actions WHERE attempt_id=? AND status='succeeded'", (attempt_id,)).fetchone() and status in ('failed', 'no_match', 'skipped'):
            raise ValueError('已有成功行动，不能改为失败以重新投递')
        if status not in ('in_progress', 'failed', 'needs_user', 'uncertain', 'no_match', 'skipped'):
            raise ValueError('成功状态只能通过核验行动结果设置')
        with self.conn:
            self.conn.execute('UPDATE attempts SET status=?, reason=?, updated=? WHERE id=?',
                              (status, reason, now(), attempt_id))

    def owns_attempt(self, run_id, attempt_id):
        attempt = self.one('attempts', attempt_id)
        return attempt['run_id'] == run_id or bool(self.conn.execute(
            'SELECT 1 FROM run_attempts WHERE run_id=? AND attempt_id=?', (run_id, attempt_id)).fetchone())

    def attach_attempt(self, run_id, attempt_id):
        config = self.check_run(run_id)
        attempt = self.one('attempts', attempt_id)
        scope = config.get('companies' if attempt['channel'] == 'web' else 'boss_companies')
        if attempt['channel'] not in config['channels'] or (scope and attempt['company'] not in scope):
            raise ValueError('历史申请不在本轮授权范围内')
        with self.conn:
            self.conn.execute('INSERT OR IGNORE INTO run_attempts VALUES (?,?,?)', (run_id, attempt_id, now()))
        return attempt

    def record_field(self, attempt_id, name, value, source, run_id=None):
        attempt = self.one('attempts', attempt_id)
        run_id = run_id or attempt['run_id']
        if not self.owns_attempt(run_id, attempt_id):
            raise ValueError('填写目标不属于当前轮次')
        self.check_run(run_id, 'fill')
        if not name or not source:
            raise ValueError('字段名和资料出处不能为空')
        with self.conn:
            self.conn.execute('''INSERT INTO fields VALUES (?,?,?,?,NULL)
              ON CONFLICT(attempt_id,name) DO UPDATE SET value=excluded.value,source=excluded.source,evidence_id=NULL''',
                              (attempt_id, name, str(value), source))

    def evidence(self, attempt_id, source, kind, fields=None):
        self.one('attempts', attempt_id)
        from PIL import Image
        source = Path(source).resolve()
        with Image.open(source) as picture:
            if picture.format != 'PNG':
                raise ValueError('证据必须为原始 PNG 截图')
            picture.verify()
        for name in fields or []:
            if not self.conn.execute('SELECT 1 FROM fields WHERE attempt_id=? AND name=?', (attempt_id, name)).fetchone():
                raise ValueError(f'未登记字段：{name}')
        identity = uuid.uuid4().hex
        dest = self.root / 'evidence' / attempt_id / f'{identity}.png'
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, dest)
        sha = hashlib.sha256(dest.read_bytes()).hexdigest()
        with self.conn:
            self.conn.execute('INSERT INTO evidence VALUES (?,?,?,?,?,?)',
                              (identity, attempt_id, kind, str(dest), sha, now()))
            for name in fields or []:
                self.conn.execute('UPDATE fields SET evidence_id=? WHERE attempt_id=? AND name=?',
                                  (identity, attempt_id, name))
        return identity

    def check_evidence(self, evidence_id, attempt_id):
        item = self.one('evidence', evidence_id)
        path = Path(item['path'])
        if item['attempt_id'] != attempt_id or not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != item['sha256']:
            raise ValueError('证据丢失、变更或不属于本次申请')
        return item

    def handoff(self, conversation, reason):
        with self.conn:
            self.conn.execute('''INSERT INTO conversations VALUES (?,?,?,?,?) ON CONFLICT(id)
              DO UPDATE SET status='human',reason=excluded.reason,updated=excluded.updated''',
                              (conversation, 'human', reason, '', now()))
        return self.attention(conversation, reason)

    def resume_conversation(self, conversation, confirmation):
        if not confirmation.strip():
            raise ValueError('恢复会话需要用户明确指令')
        self.one('conversations', conversation)
        with self.conn:
            self.conn.execute("UPDATE conversations SET status='auto', reason='',updated=? WHERE id=?", (now(), conversation))
            self.conn.execute('UPDATE attention SET resolved=1 WHERE subject=?', (conversation,))

    def observe_message(self, conversation, message_key, text):
        if not message_key or not text.strip():
            raise ValueError('消息标识和正文不能为空')
        with self.conn:
            old = self.conn.execute('SELECT text FROM messages WHERE conversation_id=? AND message_key=?',
                                    (conversation, message_key)).fetchone()
            if old:
                if old[0] != text:
                    raise ValueError('消息标识冲突，请使用可区分消息的时间与序号')
                return False
            self.conn.execute('INSERT INTO messages VALUES (?,?,?,?)', (conversation, message_key, text, now()))
            self.conn.execute('''INSERT INTO conversations VALUES (?,?,?,?,?) ON CONFLICT(id)
               DO UPDATE SET latest_message=excluded.latest_message,updated=excluded.updated''',
                              (conversation, 'auto', '', message_key, now()))
        return True

    def attention(self, subject, reason):
        old = self.conn.execute('SELECT id FROM attention WHERE subject=? AND reason=? AND resolved=0', (subject, reason)).fetchone()
        if old:
            return {'id': old[0], 'new': False}
        identity = uuid.uuid4().hex
        with self.conn:
            self.conn.execute('INSERT INTO attention (id,subject,reason,resolved,created) VALUES (?,?,?,?,?)', (identity, subject, reason, 0, now()))
        return {'id': identity, 'new': True}

    def claim_notification(self, attention_id):
        with self.conn:
            cursor = self.conn.execute('UPDATE attention SET notified=1 WHERE id=? AND notified=0 AND resolved=0', (attention_id,))
        return cursor.rowcount == 1

    def field_fingerprint(self, attempt_id):
        rows = [dict(row) for row in self.conn.execute('SELECT * FROM fields WHERE attempt_id=? ORDER BY name', (attempt_id,))]
        return hashlib.sha256(packed(rows).encode()).hexdigest()

    def prepare_action(self, run_id, attempt_id, kind, recipient, payload):
        with self.conn:
            self.conn.execute('BEGIN IMMEDIATE')
            return self._prepare_action(run_id, attempt_id, kind, recipient, dict(payload))

    def _prepare_action(self, run_id, attempt_id, kind, recipient, payload):
        if kind not in ('greet', 'submit', 'reply', 'resume'):
            raise ValueError('无效行动类型')
        config = self.check_run(run_id, kind)
        attempt = self.one('attempts', attempt_id)
        if not self.owns_attempt(run_id, attempt_id):
            raise ValueError('申请不属于当前运行')
        if not recipient:
            raise ValueError('接收方不能为空')
        if (kind == 'submit') != (attempt['channel'] == 'web'):
            raise ValueError('行动类型与渠道不符')
        conversation = self.conn.execute('SELECT * FROM conversations WHERE id=?', (recipient,)).fetchone()
        if conversation and conversation['status'] == 'human':
            raise ValueError('会话已转人工，请等待用户接管')
        if kind == 'greet':
            used = self.conn.execute("SELECT count(*) FROM actions WHERE run_id=? AND kind='greet' AND status!='not_done'", (run_id,)).fetchone()[0]
            if used >= config.get('max_contacts', 0):
                raise ValueError('已达到本轮沟通数量上限')
        if kind == 'submit':
            fields = self.conn.execute('SELECT evidence_id FROM fields WHERE attempt_id=?', (attempt_id,)).fetchall()
            if not fields or any(not f[0] for f in fields):
                raise ValueError('填写字段截图证据不完整')
            for field in fields:
                self.check_evidence(field[0], attempt_id)
            if self.check_evidence(payload.get('pre_submit_evidence', ''), attempt_id)['kind'] != 'pre_submit':
                raise ValueError('必须使用提交前复核截图')
            payload['field_fingerprint'] = self.field_fingerprint(attempt_id)
        if kind == 'resume':
            if payload.get('requested') is not True or not payload.get('resume_sha256'):
                raise ValueError('仅对方索要后允许发送指定简历')
            if config.get('resume_sha256') != payload['resume_sha256']:
                raise ValueError('简历与本轮授权版本不一致')
        if kind == 'reply' and (not payload.get('text', '').strip() or not payload.get('message_key')):
            raise ValueError('回复必须指定正文和对应消息')
        if kind in ('reply', 'resume') and conversation and conversation['latest_message'] != payload.get('message_key'):
            raise ValueError('消息已更新，必须重新判断后再回复')
        unique = [attempt['channel'], kind, recipient]
        if kind in ('greet', 'submit'):
            unique.append(attempt['job_key'])
        elif kind == 'reply':
            unique.append(payload['message_key'])
        else:
            unique.append(payload['resume_sha256'])
        key = hashlib.sha256(packed(unique).encode()).hexdigest()
        if self.conn.execute("SELECT 1 FROM actions WHERE dedup_key=? AND status!='not_done'", (key,)).fetchone():
            raise ValueError('重复行动或结果待核实，禁止重发')
        identity, stamp = uuid.uuid4().hex, now()
        with self.conn:
            self.conn.execute('INSERT INTO actions VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                              (identity, run_id, attempt_id, kind, recipient, packed(payload), key, 'pending', None, '', stamp, stamp))
        return identity

    def claim_action(self, run_id, action_id):
        with self.conn:
            self.conn.execute('BEGIN IMMEDIATE')
            action = self.one('actions', action_id)
            if action['run_id'] != run_id or action['status'] != 'pending':
                raise ValueError('行动不属于当前轮次或已执行，先核实结果')
            self.check_run(run_id, action['kind'])
            payload = json.loads(action['payload'])
            if action['kind'] == 'submit':
                if payload.get('field_fingerprint') != self.field_fingerprint(action['attempt_id']):
                    raise ValueError('填写字段已变化，必须取消未执行行动并重新复核截图')
                self.check_evidence(payload['pre_submit_evidence'], action['attempt_id'])
                for field in self.conn.execute('SELECT evidence_id FROM fields WHERE attempt_id=?', (action['attempt_id'],)):
                    self.check_evidence(field[0], action['attempt_id'])
            if action['kind'] in ('reply', 'resume'):
                state = self.one('conversations', action['recipient'])
                if state['status'] == 'human' or state['latest_message'] != payload['message_key']:
                    raise ValueError('会话转人工或消息已变化，必须重新判断')
            self.conn.execute("UPDATE actions SET status='executing',updated=? WHERE id=?", (now(), action_id))
        return action

    def finish_action(self, action_id, status, evidence_id=None, reason=''):
        action = self.one('actions', action_id)
        if status not in ('succeeded', 'uncertain', 'not_done'):
            raise ValueError('结果只能为 succeeded、uncertain 或 not_done')
        if action['status'] == 'succeeded':
            raise ValueError('已成功的行动不能覆盖')
        if status in ('succeeded', 'not_done'):
            if not reason.strip():
                raise ValueError('必须说明页面上核验到的凭据')
            evidence = self.check_evidence(evidence_id or '', action['attempt_id'])
            if evidence['kind'] not in ('result', 'failure', 'verification'):
                raise ValueError('必须使用结果或核查截图，不能用提交前图片证明成功')
        with self.conn:
            self.conn.execute('UPDATE actions SET status=?,evidence_id=?,reason=?,updated=? WHERE id=?',
                              (status, evidence_id, reason, now(), action_id))
            if action['kind'] in ('submit', 'greet') or status == 'uncertain':
                state = {'succeeded': 'success', 'uncertain': 'uncertain', 'not_done': 'in_progress'}[status]
                self.conn.execute('UPDATE attempts SET status=?,reason=?,updated=? WHERE id=?',
                                  (state, reason, now(), action['attempt_id']))

    def snapshot(self):
        return {table: [dict(r) for r in self.conn.execute(f'SELECT * FROM {table}')]
                for table in ('runs', 'attempts', 'fields', 'actions', 'conversations', 'attention', 'evidence')}

    def export(self, destination=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        destination = Path(destination or self.root / 'results.xlsx').resolve()
        book = Workbook()
        book.remove(book.active)
        data = self.snapshot()
        tabs = [('网申记录', [r for r in data['attempts'] if r['channel'] == 'web']),
                ('BOSS记录', [r for r in data['attempts'] if r['channel'] == 'boss']),
                ('待人工处理', [r for r in data['attention'] if not r['resolved']]),
                ('填写内容', data['fields']), ('行动记录', data['actions'])]
        for title, records in tabs:
            page = book.create_sheet(title)
            if not records:
                page.append(['暂无记录'])
                continue
            columns = list(records[0])
            if title in ('网申记录', 'BOSS记录'):
                columns.append('截图目录')
            page.append(columns)
            for record in records:
                values = [record.get(k, '') for k in columns]
                if columns[-1] == '截图目录':
                    values[-1] = str(self.root / 'evidence' / record['id'])
                page.append(values)
                for cell in page[page.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # 外部内容不得被 Excel 当作公式执行。
                if columns[-1] == '截图目录':
                    page.cell(page.max_row, len(columns)).hyperlink = Path(values[-1]).as_uri()
            page.freeze_panes = 'A2'
            page.auto_filter.ref = page.dimensions
            for cell in page[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='245B78')
            for col in page.columns:
                page.column_dimensions[col[0].column_letter].width = min(60, max(18, len(str(col[0].value)) * 2))
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(f'.{destination.stem}-{uuid.uuid4().hex}.xlsx')
        book.save(temporary)
        book.close()
        try:
            temporary.replace(destination)
        except PermissionError:
            raise ValueError(f'结果表被占用；最新报告已保留：{temporary}') from None
        return str(destination)
