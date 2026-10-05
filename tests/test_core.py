"""覆盖重复提交、证据缺失、越界授权等会导致真实误操作的行为。"""
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/job-search-assistant/scripts'
sys.path.insert(0, str(SCRIPTS))


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('job_assistant'), '求职运行模块尚未实现')
        from job_assistant.store import Store
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = Store(self.root)
        self.addCleanup(self.db.close)

    def test_attempt_survives_reopen_and_duplicate_is_blocked(self):
        run = self.db.create_run({'channels': ['web'], 'max_contacts': 2}, '用户确认本轮')
        first = self.db.begin_attempt(run, '甲公司', 'https://example.com/job/1', '嵌入式', 'web', 'job-1')
        self.db.set_attempt(first, 'uncertain', '提交后超时')
        self.db.close()
        from job_assistant.store import Store
        self.db = Store(self.root)
        self.addCleanup(self.db.close)
        with self.assertRaisesRegex(ValueError, '重复|核实'):
            self.db.begin_attempt(run, '甲公司', 'https://example.com/job/1', '嵌入式', 'web', 'job-1')

    def test_web_company_table_does_not_restrict_boss_recommendations(self):
        run = self.db.create_run({'channels': ['web', 'boss'], 'companies': ['甲公司'], 'max_contacts': 2}, '确认两个渠道')
        attempt = self.db.begin_attempt(run, '乙公司', 'https://www.zhipin.com/job/2', '嵌入式', 'boss', '2')
        self.assertEqual(self.db.one('attempts', attempt)['company'], '乙公司')
        with self.assertRaisesRegex(ValueError, '范围'):
            self.db.begin_attempt(run, '乙公司', 'https://example.com/job/2', '嵌入式', 'web', '2')

    def test_submission_needs_field_coverage_and_real_png(self):
        run = self.db.create_run({'channels': ['web'], 'max_contacts': 2}, '用户确认本轮')
        attempt = self.db.begin_attempt(run, '甲公司', 'https://example.com', '嵌入式', 'web', '1')
        self.db.record_field(attempt, '姓名', '张三', 'profile.name')
        with self.assertRaisesRegex(ValueError, '截图|证据'):
            self.db.prepare_action(run, attempt, 'submit', '甲公司', {'url': 'https://example.com'})

    def test_complete_evidence_allows_submit_and_mutated_field_invalidates_coverage(self):
        from PIL import Image
        run = self.db.create_run({'channels': ['web']}, '用户确认')
        a = self.db.begin_attempt(run, '甲', 'https://example.com', '嵌入式', 'web', '1')
        self.db.record_field(a, '姓名', '张三', 'profile.name')
        picture = self.root / 'screen.png'
        Image.new('RGB', (100, 100), 'white').save(picture)
        evidence = self.db.evidence(a, picture, 'pre_submit', ['姓名'])
        action = self.db.prepare_action(run, a, 'submit', '甲', {'pre_submit_evidence': evidence})
        receipt = self.db.evidence(a, picture, 'result')
        self.db.finish_action(action, 'succeeded', receipt, '模拟网站回执 TEST-1')
        self.assertEqual(self.db.one('attempts', a)['status'], 'success')
        output = self.db.export()
        from openpyxl import load_workbook
        book = load_workbook(output)
        self.assertIn('填写内容', book.sheetnames)
        book.close()

    def test_changed_field_after_preparation_blocks_commit(self):
        from PIL import Image
        run = self.db.create_run({'channels': ['web']}, '用户确认')
        a = self.db.begin_attempt(run, '甲', 'https://example.com', '工程师', 'web', '1')
        self.db.record_field(a, '姓名', '张三', 'name')
        pic = self.root / 'screen.png'
        Image.new('RGB', (100, 100)).save(pic)
        evidence = self.db.evidence(a, pic, 'pre_submit', ['姓名'])
        action = self.db.prepare_action(run, a, 'submit', '甲', {'pre_submit_evidence': evidence})
        self.db.record_field(a, '姓名', '李四', 'name')
        self.assertTrue(hasattr(self.db, 'claim_action'), '缺少执行前复核')
        with self.assertRaises(ValueError):
            self.db.claim_action(run, action)

    def test_success_cannot_be_downgraded_to_enable_duplicate(self):
        from PIL import Image
        run = self.db.create_run({'channels': ['boss'], 'max_contacts': 2}, '用户确认')
        a = self.db.begin_attempt(run, '甲', 'https://www.zhipin.com', '工程师', 'boss', '1')
        action = self.db.prepare_action(run, a, 'greet', 'hr', {})
        pic = self.root / 'screen.png'
        Image.new('RGB', (100, 100)).save(pic)
        evidence = self.db.evidence(a, pic, 'result')
        self.db.finish_action(action, 'succeeded', evidence, '出现首次沟通消息')
        with self.assertRaises(ValueError):
            self.db.set_attempt(a, 'failed', '想重新投递')

    def test_pre_submit_image_cannot_prove_success(self):
        from PIL import Image
        run = self.db.create_run({'channels': ['boss'], 'max_contacts': 2}, '用户确认')
        a = self.db.begin_attempt(run, '甲', 'https://www.zhipin.com', '工程师', 'boss', '1')
        action = self.db.prepare_action(run, a, 'greet', 'hr', {})
        pic = self.root / 'screen.png'
        Image.new('RGB', (100, 100)).save(pic)
        evidence = self.db.evidence(a, pic, 'pre_submit')
        with self.assertRaises(ValueError):
            self.db.finish_action(action, 'succeeded', evidence, '没有实际回执')

    def test_attention_can_be_notified_once_after_handoff(self):
        item = self.db.handoff('hr', 'interview')
        self.assertTrue(hasattr(self.db, 'claim_notification'), '缺少通知去重记录')
        self.assertTrue(self.db.claim_notification(item['id']))
        self.assertFalse(self.db.claim_notification(item['id']))

    def test_latest_message_cannot_be_replaced_by_old_observation(self):
        self.db.observe_message('hr', 'm1', '你好')
        self.db.observe_message('hr', 'm2', '明天来面试')
        self.assertFalse(self.db.observe_message('hr', 'm1', '你好'))
        self.assertEqual(self.db.one('conversations', 'hr')['latest_message'], 'm2')

    def test_formula_like_external_text_is_exported_as_text(self):
        run = self.db.create_run({'channels': ['web']}, '用户确认')
        self.db.begin_attempt(run, '=HYPERLINK("bad")', 'https://example.com', '工程师', 'web', '1')
        from openpyxl import load_workbook
        book = load_workbook(self.db.export())
        try:
            page = book['网申记录']
            index = [c.value for c in page[1]].index('company') + 1
            self.assertEqual(page.cell(2, index).data_type, 's')
        finally:
            book.close()

    def test_pause_and_stop_block_new_external_actions(self):
        run = self.db.create_run({'channels': ['boss'], 'max_contacts': 2}, '用户确认本轮')
        attempt = self.db.begin_attempt(run, '甲公司', 'https://www.zhipin.com', '嵌入式', 'boss', '1')
        self.db.control(run, 'paused')
        with self.assertRaisesRegex(ValueError, '暂停|运行'):
            self.db.prepare_action(run, attempt, 'greet', 'hr1', {})
        self.db.control(run, 'stopped')
        with self.assertRaises(ValueError):
            self.db.control(run, 'running')

    def test_contact_limit_counts_uncertain_actions(self):
        run = self.db.create_run({'channels': ['boss'], 'max_contacts': 1}, '用户确认本轮')
        a = self.db.begin_attempt(run, '甲', 'https://www.zhipin.com/1', '嵌入式', 'boss', '1')
        self.db.prepare_action(run, a, 'greet', 'hr1', {})
        b = self.db.begin_attempt(run, '乙', 'https://www.zhipin.com/2', '嵌入式', 'boss', '2')
        with self.assertRaisesRegex(ValueError, '上限'):
            self.db.prepare_action(run, b, 'greet', 'hr2', {})

    def test_exact_message_is_not_sent_twice_and_handoff_blocks_reply(self):
        run = self.db.create_run({'channels': ['boss'], 'max_contacts': 2}, '用户确认本轮')
        a = self.db.begin_attempt(run, '甲', 'https://www.zhipin.com/1', '嵌入式', 'boss', '1')
        self.db.prepare_action(run, a, 'reply', 'hr1', {'text': '您好', 'message_key': 'msg1'})
        with self.assertRaisesRegex(ValueError, '重复|核实'):
            self.db.prepare_action(run, a, 'reply', 'hr1', {'text': '您好', 'message_key': 'msg1'})
        self.db.handoff('hr1', '面试邀约')
        with self.assertRaisesRegex(ValueError, '人工'):
            self.db.prepare_action(run, a, 'reply', 'hr1', {'text': '收到', 'message_key': 'msg2'})


class InputTests(unittest.TestCase):
    def test_csv_chinese_headers_urls_and_duplicate_rows(self):
        self.assertIsNotNone(importlib.util.find_spec('job_assistant'), '输入模块尚未实现')
        from job_assistant.inputs import read_companies
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '公司.csv'
            path.write_text('公司名称,申请网站\n甲,https://example.com\n甲,https://example.com\n', encoding='utf-8-sig')
            rows = read_companies(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]['company'], '甲')
            path.write_text('公司名称,申请网站\n甲,javascript:alert(1)\n', encoding='utf-8-sig')
            with self.assertRaises(ValueError):
                read_companies(path)


if __name__ == '__main__':
    unittest.main()
